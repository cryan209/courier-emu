from pathlib import Path
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path('/Users/scottcryan/courier-emu')
OUT=ROOT/'artifacts/x2-spec-20261004'
d=Document(); sec=d.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(21);sec.bottom_margin=Mm(20);sec.left_margin=Mm(22);sec.right_margin=Mm(22)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(11 if name=='Normal' else 14 if name=='Heading 1' else 12 if name=='Heading 2' else 10 if name=='Caption' else 22 if name=='Title' else 12)
 s.paragraph_format.space_after=Pt(6)
 s.paragraph_format.line_spacing=1.08
 for x in s.element.xpath('.//w:rFonts'):
  for key in list(x.attrib):
   if 'theme' in key.lower():del x.attrib[key]
 for x in s.element.xpath('.//w:pBdr'):x.getparent().remove(x)
for name in ['Heading 1','Heading 2','Caption']:d.styles[name].paragraph_format.keep_with_next=True
d.styles['Title'].font.bold=True
h=sec.header.paragraphs[0];h.text='COURIER EMU     x2 firmware reconstruction     Draft 0.3';h.runs[0].font.size=Pt(9)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('x2 technical specification     •     4 October 2026     •     ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for r in f.runs:r.font.size=Pt(9)
def p(t,style=None):return d.add_paragraph(t,style)
def hd(t):d.add_heading(t,1)
def sub(t):d.add_heading(t,2)
def page():d.add_page_break()
def table(cap,heads,rows,widths):
 p(cap,'Caption');t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Mm(w)
 for c,x in zip(t.rows[0].cells,heads):c.text=x
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for n,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  for c,w in zip(row.cells,widths):
   c.width=Mm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b)
   if n==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');margins.append(x)
   pr.append(margins)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1
    for r in pp.runs:r.font.size=Pt(10);r.bold=(n==0)
 p('')
def code(t):
 for line in t.splitlines():
  pp=p(line);pp.paragraph_format.space_after=Pt(0)
  for r in pp.runs:r.font.name='Courier New';r.font.size=Pt(9)

p('x2 Modem Technical Specification','Title')
p('Negotiation and downstream mapping','Subtitle')
p('Draft 0.3   •   Independent firmware reconstruction   •   4 October 2026')
hd('1 Scope and operating model')
p('This document explains how the recovered x2 implementation selects an asymmetric connection, announces directional modulation parameters, starts its training source and prepares downstream signal tables. It covers the Courier 3453C analogue client, the original I-modem server and additional mapper evidence from the later Quad server firmware.')
p('The client retains executable x2 setup and DSP code. The server capability exchange and sample transforms have been reconstructed. The later Quad reconstruction now connects a received parameter record to six mapping positions and the expansion of candidate codeword sets into six working banks. The stock controller validates the overlay source placement. Startup construction and its consumers also execute together in bounded test contexts. Payload mapping and complete peer negotiation remain open.')
p('The digital server supplies downstream sample codewords. The analogue client receives the resulting waveform and must recover the information represented by the chosen levels. The upstream direction uses the related V.34 machinery. Negotiation therefore has to agree both the directional modulation choice and the parameters that make the downstream levels distinguishable at the client.')
p('A constellation is the set of signal levels available for carrying data. A mapper selects members of that set from input data; the receiver performs the inverse operation. A table of levels is only one part of that process. It does not, on its own, establish data framing, shaping, synchronization or payload decoding.')
hd('2 Status of the specification')
p('The numbered clauses and field tables follow the style of an ITU technical specification. This is an independent engineering draft, not an ITU-T Recommendation. Statements describe the recovered firmware profile; they are not universal requirements for all x2 equipment.')
table('Table 1 — Evidence used in this draft',['Designation','Meaning'],[('Verified','A semantic reconstruction matches original DSP instructions for the tested inputs and entry conditions.'),('Observed','Firmware instructions, constants or message construction establish the described operation; a complete call has not been validated.'),('Provisional','Component behavior is reproducible, but a placement assumption or protocol interpretation remains unverified.'),('Open','The available evidence does not yet establish the procedure.')],[30,136])
p('Firmware addresses, memory layouts and reproduction commands are confined to Annex A. The main clauses describe connection behavior and signal parameters. Timing and interoperability requirements are stated only where the evidence supports them.')

page();hd('3 Negotiation sequence')
p('Table 2 organizes the recovered operations into a connection sequence. The individual operations are supported to different degrees; their complete ordering and transition timing have not been demonstrated in one end-to-end call.')
table('Table 2 — Recovered negotiation stages',['Stage','Purpose and observed behavior','Status'],[('Local selection','Check whether x2 is enabled and permitted by the selected mode, speed settings and line conditions. Prepare its capability announcement.','Observed'),('Capability exchange','Use the INFO0 capability structure to announce permitted asymmetry and related attributes. The server route into the outgoing capability buffer is verified.','Verified in server'),('Directional announcement','Transmit a short proprietary body naming a carrier choice and two modulation-rate indices. One index can identify the PCM direction.','Observed'),('Channel qualification','Evaluate peer parameters and a signed channel margin. Related rate machinery limits the requested rate against available capability.','Partial'),('Training and probing','Generate a repeating startup source and collect received training words and a position bitmap.','Source verified'),('Constellation preparation','Decode received choices, select candidate level sets and expand six position banks.','Components verified'),('Data operation','Map payload to downstream codewords and recover payload at the client. The complete inverse and frame-alignment procedure remain open.','Open')],[34,103,29])
sub('3 1 Local selection and fallback')
p('The client calls a distinct x2 eligibility procedure during connection setup. It excludes disabled x2 operation, unsuitable modem modes and speed or line conditions that fail the firmware gates. The neighboring V.90 selection procedure is separate. Retaining an x2 setup call establishes an active selection mechanism, but does not prove successful peer negotiation.')
p('The setup prepares a capability pattern describing directional asymmetry. Some settings can suppress that asymmetry. A later announcement can name equal modulation indices in both directions, providing evidence of a symmetric alternative. The precise transition from rejected x2 operation to an established fallback connection is still open.')
sub('3 2 What remains to connect')
p('The client stores the x2 capability supplied by its supervisor, but the consumer that connects that particular store to the transmitted capability buffer has not been traced. Its tested initialization fragment reads a different capability word. This missing connection is specific to the 3453C profile and must be resolved before claiming complete client negotiation.')

page();hd('4 Capability and modulation announcements')
sub('4 1 INFO0 capability fields')
p('The interpreted client capability announcement uses fields already present in the V.34 INFO0 structure. The supervisor sets the permitted asymmetry, clears the CME indication and clears the clock-source field. Table 3 states the interpreted wire field correspondence. The server capability word has a verified route into its INFO0 work buffer; actual frame emission and peer acceptance remain separate checks.')
table('Table 3 — Interpreted capability fields',['INFO0 bits','Parameter','Client construction'],[('21 to 23','Maximum directional symbol-rate difference','An integer from 0 to 3, selected by the two relevant S58 settings'),('24','CME indication','Cleared'),('26 to 27','Transmit clock source','Cleared')],[29,63,74])
p('With both relevant S58 options clear, the asymmetry integer is 2. Enabling the first option alone makes it 3; enabling the second alone makes it 0; enabling both makes it 1. These are numerical values, not printed bit strings that might be confused with transmission order.')
sub('4 2 Directional modulation parameters')
p('The proprietary body has seven bits. The lowest bit selects the carrier; the next three bits hold the first modulation index and the highest three hold the second. The roles of the two indices interchange with the modem role flag. Index 6 is interpreted as the PCM-direction indication; ordinary V.34 modulation indices occupy the range 0 to 5.')
table('Table 4 — Directional announcement body',['Body bits','Width','Meaning'],[('0','1','Carrier choice: 0 low, 1 high'),('1 to 3','3','First modulation index'),('4 to 6','3','Second modulation index')],[28,20,118])
table('Table 5 — Observed announcement values',['Body value','Carrier','First index','Second index'],[('4D hexadecimal','High','6','4'),('69 hexadecimal','High','4','6'),('48 hexadecimal','Low','4','4'),('49 hexadecimal','High','4','4')],[52,38,38,38])
p('The first two values announce the same asymmetric pair with its directions exchanged. The latter two name equal indices and differ only in carrier choice. These values come from the inspected Courier announcement builders; they have not been established as a complete set of every possible x2 peer message.')

page();hd('5 Framing recognition and rate agreement')
sub('5 1 Transmitted announcement structure')
p('The recovered x2-only announcement script contains two lead-in ones, twelve bits of fill and synchronization, the seven-bit body, sixteen check bits and four closing ones. The framed portion is 39 bits; including the lead-in gives 41 emitted script bits. A lead-in is therefore not counted as part of the 39-bit frame.')
table('Table 6 — Announcement script in emission order',['Element','Bits','Interpretation'],[('Lead-in','2','Two ones before the framed portion'),('Fill and synchronization','12','Shared INFO framing sequence'),('Information body','7','Carrier and directional modulation indices'),('Check field','16','INFO CRC family'),('Closing fill','4','Four ones')],[60,20,86])
p('The local field encoder models the body least-significant bit first and the V.34 INFO CRC initialized to all ones, using the reflected form of the polynomial with terms of order 16, 12, 5 and 0. This is a working serialization model. A captured x2 frame or an executed complete serializer is still needed to validate every bit convention.')
sub('5 2 Receiver recognition remains open')
p('A seven-bit comparison in the shared receiver was previously interpreted as recognition of the proprietary announcement. That interpretation is superseded: the same comparison exists in pre-x2 V.34 firmware and matches a window of the INFO fill and synchronization preamble. It does not identify an x2 body value.')
p('No receiver for the proprietary seven-bit body has been conclusively identified. The framing script proves construction of an announcement; it does not prove how a server accepts it. The document therefore does not assign an acknowledgement, timeout or retry procedure to this message.')
sub('5 3 Channel qualification and rate limits')
p('The 3453C qualification fragment tests a peer attribute and computes a signed margin from peer parameters and a channel measurement. A positive margin enables a subsequent state choice. The physical unit of that margin and the peer identity implied by the test remain open.')
p('Related later firmware compares a host-requested maximum with the existing maximum and retains the smaller value. Its two directional rate fields exchange roles with the same role flag used by the directional announcement. It also intersects enabled rate capabilities with a symbol-rate-dependent mask and the extended-constellation capability. This is evidence of inherited V.34 rate agreement; its complete x2-specific binding remains open.')
p('Repository rate analysis supports a six-symbol downstream organization with nominal rate increments of 8,000 ÷ 6 bit/s. The new six-position mapper evidence is consistent with that organization. Consistency does not establish its alignment signal, negotiated bit allocation or exact rate-selection procedure.')

page();hd('6 Startup and receiver qualification')
sub('6 1 Repeating startup source')
p('The original server has a four-state startup source. Starting from its initialized state, it repeats the 2,010-call cycle in Table 7. Execution comparison over 4,096 consecutive source calls covers two complete cycles and part of a third.')
table('Table 7 — Verified server source cycle',['Part','Source words','Count'],[('Training prefix','Repeated hexadecimal 7E','1,747'),('Separator','Zero','7'),('Probe pairs','Ascending first word and descending second word: (0,255), (1,254), continuing to (127,128)','128 pairs'),('Restart','Return to the training prefix','Cycle repeats')],[34,102,30])
p('The pair generator advances the ascending value after selecting its return word, then emits the descending value before decrementing it. On the last pair, it restores the prefix count and returns to the first state. The reconstructed delayed-return updates match the original instructions.')
p('These are source words. The selected output path may mask their width, scramble them and reverse an octet before transmission. A source-call count does not establish elapsed time. The sample cadence and callback selected during each startup stage must be traced before specifying an on-wire duration.')
sub('6 2 Received training and position bitmap')
p('The inspected receive chain begins by qualifying repeated training values. It advances a position counter over six positions and accumulates a bitmap. Later code classifies received low-valued patterns, records words in a small report buffer and changes receive substates. The initial detector seeks seven qualifying training values before advancing.')
p('The bitmap and recorded reports are plausible inputs to channel qualification and constellation selection, but their full relationship to the outgoing probe pairs is not established. The receiver has not yet been execution compared as a complete sequence. Its report format, rejection conditions and transition deadlines remain open.')
sub('6 3 Transition to operating transforms')
p('The server sample path can select seven-bit or eight-bit source words. Transmission applies the selected mask, optional scrambling, optional octet reversal and an output write. Reception records a sample, performs the inverse reversal, applies the mask and optionally descrambles it before delivery.')
p('The tested sample transforms match the original routines with payload source and sink dispatch disabled. They establish transformation behavior, but not supervisor queue delivery or a complete data connection. A seven-bit digital bearer word must not be interpreted as the downstream analogue x2 rate without the mapper and frame allocation.')

page();hd('7 Negotiated downstream parameters')
sub('7 1 Transfer of the received record')
p('In the inspected PCM receive branch, four buffered words become the working negotiation record. The first two words supply directional parameters; the last two supply the position choices and an additional parameter. This transfer and the subsequent unpacking match the original instructions. The branch condition is part of the verified entry context; the check does not include frame reception, CRC acceptance or acknowledgement.')
table('Table 8 — Recovered working parameters',['Parameter','Recovered representation','Qualification'],[('Original mode','Four bits from the first working word','Used to select a scale table'),('Effective mode','Original mode, or zero when a local override is set','Used by a second parameter table'),('Additional parameter','Three bits above the position record','Physical meaning remains open'),('Position record','Twenty-four bits organized as six four-bit cells','Three bits of each cell select a control value')],[37,71,58])
p('The distinction between original and effective mode is observable: a local override changes the stored mode to zero, while one scale-table lookup retains the original mode. The draft therefore keeps both parameters distinct. Their exact rate names and the physical units of the scale values have not been established.')
sub('7 2 Conversion of position choices')
p('The firmware consumes the six cells from the most significant end of its working record. It ignores the upper bit of each cell and converts the lower three bits according to Table 9. Each result occupies two bits in the position-control word. Bank expansion then consumes that word from its least significant end, so its first bank corresponds to the last cell processed during unpacking. This describes working-record order; wire order remains to be established.')
table('Table 9 — Cell value to position control',['Cell value','0','1','2','3','4','5','6','7'],[('Control','0','1','2','1','0','3','0','3')],[38,16,16,16,16,16,16,16,16])
p('The firmware also counts the nonzero controls. If all six controls are zero, it substitutes a control word with the lowest position set to value 2, while retaining an active-position count of zero. This is a distinct default state. It must not be described as six ordinary zero choices or silently recounted as one active position.')
sub('7 3 Extent of the verified negotiation path')
p('Tests cover received-record transfer, parameter unpacking and expansion into six banks for all 4,096 possible control combinations, with both values of the ignored cell bit represented. Candidate sets are supplied before expansion in these tests. Their construction between unpacking and expansion has not yet been execution compared as one negotiated call. The relation to the received training bitmap, agreed bit rate and transmitted parameter frame remains open.')

page();hd('8 Downstream mapping organization')
sub('8 1 Six positions and two candidate sets')
p('Six two-bit controls select the candidate sets for the six positions. A size-distribution routine assigns one of two set sizes to each position: control zero chooses the high-byte size and controls 1, 2 or 3 choose the low-byte size. The count and size routines match the original instructions for all 4,096 combinations.')
table('Table 10 — Position control behavior',['Control','Candidate set','Additional expansion flag'],[('0','First candidate','Clear'),('1','Second candidate','Set'),('2','Second candidate','Clear'),('3','Second candidate','Set')],[25,75,66])
p('The expansion flag is a separate stored attribute. Its interpretation as sign, orientation or another signal property is not yet established. The firmware preserves it as part of each working entry.')
sub('8 2 Expansion into working banks')
p('The expansion component constructs six working banks from the two candidate tables. It consumes one control per bank, chooses the corresponding source set and copies its configured length. An odd control sets the additional flag in each copied entry. Successive banks begin 128 words apart; that spacing is a storage stride, not a requirement to use 128 levels.')
p('The expansion uses a last-index convention: a configured value of three copies four entries. The startup caller supplies the value nine, so expansion copies ten entries, whereas the constructor and level-conversion loop each prepare nine. Tests seed the additional entry explicitly. Its initialization and purpose in the complete connection still require a caller trace.')
sub('8 3 Constructor and consumer execution')
p('The original startup caller has now been executed through candidate construction, level conversion and six-bank expansion in 128 seeded contexts. It sets all six size slots to nine, preserves the selected constructor codewords in the low octets, and copies the transformed candidate entries into the banks selected by the controls. These checks cover the unmodified caller and its callees together.')
p('The second candidate table and additional entries are seeded in those contexts. The result verifies their consumption and the constructor relationship; it does not establish how both candidates are populated in an actual negotiated connection. The intermediate level converter has been executed, but its amplitude and companding interpretation has not been independently reconstructed.')
sub('8 4 Verified source placement')
p('A fresh stock controller run copied all seven overlay sources into each of the four modem RAM banks. All 28 byte comparisons matched the corrected source offsets over the complete overlay lengths. The earlier one-word discrepancy was a flat-file extraction error. The downloader does not discard a leading word. This removes the placement qualification from the component results in clauses 7 to 9.')

page();hd('9 Codeword sets and signal realization')
sub('9 1 Paired constructor tables')
p('The constructor selects one of the two nine-entry tables in Table 11 according to a local mode flag. Every value in the second table is sixteen lower than the corresponding value in the first. These are alternative constructor tables; they must not be equated automatically with the two candidates consumed by bank expansion. The corrected source placement and original startup caller both support the table contents.')
table('Table 11 — Paired constructor values',['Entry','First set','Second set'],[(i+1,a,b) for i,(a,b) in enumerate(zip(['A5','A7','AD','AF','B7','BD','C5','CF','E5'],['95','97','9D','9F','A7','AD','B5','BF','D5']))],[40,63,63])
p('All entries are hexadecimal codeword values in stored table order. The difference of sixteen is a relationship between codeword indices. It shall not be converted into a fixed amplitude ratio or assigned to a companding law until the selected law and subsequent conversion path are established.')
sub('9 2 From constellation choice to a transmitted sample')
p('The recovered preparation chain selects parameter tables, derives candidate sizes, distributes them across six positions and expands working banks. The preceding mapping block also processes six values and uses auxiliary level tables. Together these establish a concrete six-position preparation architecture, beyond an isolated list of codewords.')
p('A complete payload mapper still needs the rule that converts a group of input bits into six selected entries, any redundancy or shaping constraints, the sign and flag interpretation, the conversion from stored entries to actual PCM samples, and the exact frame boundary. The analogue receiver must then implement the corresponding inverse after channel equalization.')
sub('9 3 Remaining negotiation and mapping requirements')
p('The next required checks are to trace the received training bitmap into the parameter record, execute the negotiated candidate builder with its consumers, establish the additional bank entry, and connect the agreed rate to the input-bit grouping. A full call trace must establish the alignment signal and show that the client accepts the selected constellations.')
p('No interoperability claim is made from the present component results. The document now specifies the recovered preparation behavior and identifies precisely which procedures still prevent an end-to-end x2 implementation.')

page();hd('Annex A Firmware evidence and reproduction')
p('This annex is informative. It retains implementation detail for traceability without making addresses part of the protocol description. The manually reconstructed C is not vendor source.')
table('Table A 1 — Firmware profiles',['Profile','Image','Role'],[('Client','3453C version 2.3.33 XMF','Retained x2 setup, low overlay and PCM training'),('Original server','IM020104.NAC','Capability packing, startup source and sample transforms'),('Later server','QF060003.NAC','Shared PCM helpers and six-position parameter decoding')],[30,59,77])
p('Client evidence is in artifacts/3453c-x2-decomp-20261004. Server evidence is in artifacts/x2-server-decomp-20261004. Mapper and placement evidence is in artifacts/x2-mapper-negotiation-20261004, including mapper_lift.c, mapper-alignment.asm and verification.json.')
table('Table A 2 — Comparison coverage',['Group','Cases','Qualification'],[('Client','1,769','Original mapped client instructions'),('Original server','6,474','Original mapped server instructions'),('Quad helpers','88','Corrected stock-verified source placement'),('Mapper C comparisons','13,504','Transfer, unpacking, count, sizes and bank expansion'),('Startup caller contexts','128','Unmodified constructor and consumers; seeded candidates'),('Overlay source copies','28','Seven overlays across four stock controller RAM banks')],[49,23,94])
p('Mapper entries are AB06 for received-record transfer, C870 for unpacking, C6E5 for the six-position count, C6F4 for size distribution and C92A for descriptor expansion. The first candidate source is 0A60 and the second is 0880; expanded banks begin at E8F4 with a stride of 0080. The codeword constructor is C9BF; the integrated startup caller begins at C840.')
p('Reproduce with tools/recover_3453c_x2.py, tools/recover_x2_server.py and tools/recover_x2_mapper.py using the repository virtual environment. The mapper script compiles the manual C lift and compares its data effects with original instructions in the native emulator. It also executes the original startup chain and checks its resulting banks. Run tools/verify_quad_pcm_placement.py for the independent stock controller source-copy check; its report records the corrected offsets and complete copy lengths.')
p('Negotiation construction and historical corrections are documented in docs/x2-v90-protocol-selection.md. In particular, its preamble-detector correction supersedes the earlier proposed receiver for the seven-bit body. The standard field terminology follows ITU-T V.34 Table 14; official publication: https://www.itu.int/rec/T-REC-V.34-199610-S/en. V.90 context: https://handle.itu.int/11.1002/1000/4493.')
p('Draft 0.3 adds verified source placement, received-parameter decoding, exhaustive position-control coverage and integrated startup constructor checks. It distinguishes alternative constructor tables from the two expansion candidates. Complete call negotiation and downstream payload mapping remain open.')
d.core_properties.title='x2 Modem Technical Specification'
d.core_properties.subject='Negotiation and downstream mapping firmware reconstruction'
d.core_properties.author='Courier emu project'
d.save(OUT/'x2-technical-specification.docx')
print(OUT/'x2-technical-specification.docx')
